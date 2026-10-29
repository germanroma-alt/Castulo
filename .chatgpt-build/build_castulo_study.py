#!/usr/bin/env python3
# Build an illustrated archaeological study from the repository's own files.
import os, re, sys, hashlib, textwrap, traceback
from pathlib import Path
from collections import defaultdict
from io import BytesIO

from PIL import Image as PILImage, ImageOps
from docx import Document
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, PageBreak,
    Table, TableStyle, Image, KeepTogether, HRFlowable, LongTable, Flowable)
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

ROOT = Path("OneDrive/Documents/1-ARQUEOLOGÍA/2-INDUSTRIA-ÓSEA/1-Industria-ösea-de-Cástulo/Clasificación-Huesos-de-Castulo-(7-10-26)")
OUT = Path("Estudio_extenso_industria_osea_Castulo_ilustrado.pdf")
WORK = Path("/tmp/castulo_pdf_assets")
WORK.mkdir(parents=True, exist_ok=True)

FAMILY_NAMES = {
 "A": "Hilado, tejido y costura",
 "B": "Vida cotidiana, adorno personal y toilette",
 "C": "Vestido",
 "D": "Escritura",
 "E": "Juego y divertimento",
 "F": "Elementos para muebles",
 "G": "Mundo militar",
 "H": "Herramientas y elementos varios",
 "I": "Objetos de función desconocida",
 "J": "Fases de elaboración"
}
FAMILY_TEXT = {
"A": {
 "intro": "El grupo A reúne instrumentos ligados a la preparación de fibras, al hilado, al tejido y a operaciones de costura o perforación. La documentación fotográfica reconoce ruecas y útiles de hilado, placas de telar o telarcillos, agujas y punzones. La asociación funcional es plausible, pero no permite inferir por sí sola la intensidad productiva ni el carácter doméstico o especializado del trabajo.",
 "models": "A.II comprende ruecas y elementos de hilado. En los materiales tipológicos se distinguen remates anulares y piezas torneadas que pudieron actuar como bobinas. A.V recoge placas de telar: se describen una placa rectangular con cuatro perforaciones y una placa de rejilla. A.VII corresponde a agujas (acus), incluidos los ojales en ocho obtenidos mediante perforaciones tangenciales; se diferencian remates rectos y cónicos. A.IX recoge punzones de extremos redondeados o planos. Los rasgos diagnósticos son la sección del vástago, la simetría, la posición de orificios, el acabado superficial y las huellas de fabricación.",
 "parallels": "Las fichas comparativas citan paralelos en Augusta Emerita, Asturica Augusta, Conimbriga, Villafáfila, Vareia, Torre de Palma, Baelo, Gades, Emporiae y Ammaia. Para placas de rejilla y objetos textiles se mencionan también Pompeya, South Shields, Stobi, Aquincum y Xanten. Estos lugares son paralelos tipológicos citados por la documentación; no significan que cada ejemplar de Cástulo proceda de los mismos talleres.",
 "chron": "Las cronologías propuestas para modelos comparativos son amplias: desde el siglo I d.C. hasta los siglos IV-V d.C. según el tipo. Los telarcillos de rejilla cuentan con paralelos altoimperiales, mientras que las agujas de ojal en ocho presentan una perduración prolongada. No se traslada automáticamente una fecha de paralelo al código de Cástulo sin contexto estratigráfico propio."
},
"B": {
 "intro": "El grupo B es el conjunto más numeroso del inventario digitalizado y abarca objetos de uso personal: alfileres de tocado, instrumentos de toilette, cucharillas, recipientes, contenedores, brazaletes, figurillas, colgantes y accesorios. Su amplitud explica una parte importante de su peso estadístico. La distinción entre objetos acabados, fragmentos y elementos de identificación dudosa debe mantenerse en todo el análisis.",
 "models": "B.I comprende acus crinalis o alfileres de tocado. Las fichas distinguen vástagos fusiformes y cónicos y múltiples formas de remate: esférico, semiesférico, afacetado, bulboso, barriloide, piramidal, cónico, en cáliz o en punta de oliva. B.IV reúne espátulas de toilette; B.V, cucharillas de perfume; B.VI-B.VIII, contenedores y recipientes; B.IX, strigiles; B.XII, pulseras y brazaletes; B.XV, figurillas antropomorfas; B.XVII, colgantes y cuentas; B.XIX, elementos interpretados como parte de sombrillas. Para asignar un subtipo se requieren rasgos observables, no solo la similitud del nombre de carpeta.",
 "parallels": "Las fichas nombran Augusta Emerita, Torre Águila, Conimbriga, Gades, Emporiae, Ilici, Carteia, Baelo Claudia, Asturica Augusta, Pompaelo, Pollentia y otros yacimientos peninsulares. Cástulo aparece expresamente en los paralelos de la espátula B.IV.3 y la cucharilla B.V.3. La ficha complementaria de Harpócrates se registra como colgante B.XVII.4.1.1.",
 "chron": "Los alfileres de tocado cuentan con modelos documentados entre los siglos I y V d.C. La espátula B.IV.3 se compara con contextos de los siglos III-IV d.C.; las cucharillas de perfume tienen especial presencia en los siglos I-II, aunque algunos modelos perduran más. Los colgantes y cuentas tienen cronologías variables. La ficha de Harpócrates propone I-II d.C. Todas estas dataciones corresponden a modelos o paralelos, no sustituyen la fecha contextual de cada ejemplar."
},
"C": {
 "intro": "El grupo C reúne objetos relacionados con la vestimenta. En la carpeta de imágenes se identifican cuatro códigos atribuidos a broches. El tipo se debe estudiar por su cuerpo, extremos, ranura, sección y medidas; su función concreta depende de la forma y de la relación con el vestido o con un elemento de sujeción.",
 "models": "La ficha C.II.1.1 describe un objeto alargado de sección semicilíndrica, con una ranura central rebajada y extremos abiertos en forma troncocónica que rematan en discos o semicasquetes. Se consignan medidas comparativas de aproximadamente 2-4 cm de longitud y 0,3-0,6 cm de grosor. C.II.1.2 presenta extremos hemisféricos y dimensiones menores. La denominación de archivo no permite asignar individualmente estos subtipos a los cuatro códigos sin revisar su imagen y ficha.",
 "parallels": "Cástulo aparece expresamente en la ficha comparativa del tipo C.II.1.1. Se citan también Palencia, Numancia, Augusta Emerita, Balsa, Hasta Regia, Puig des Molins, Saguntum, Emporiae y Els Antigons; fuera de Hispania figuran Nîmes, Lyon, Fréjus, Avenches y Pompeya.",
 "chron": "Los paralelos cubren desde el siglo I d.C. hasta el III d.C., con diferencias entre yacimientos y contextos. La horquilla debe tratarse como marco comparativo general y no como datación automática de los cuatro códigos de Cástulo."
},
"D": {
 "intro": "El grupo D recoge instrumentos relacionados con la escritura. La carpeta contiene un stylus, instrumento empleado habitualmente para escribir o incidir sobre superficies preparadas, como tablillas enceradas. La presencia de una pieza de esta clase es compatible con prácticas de escritura, contabilidad o anotación, pero no basta para caracterizar la alfabetización o las funciones administrativas de un sector urbano.",
 "models": "La clasificación documental disponible identifica D.I - Stylos. Para cada ejemplar habría que registrar el extremo apuntado, el extremo opuesto usado para borrar o alisar, la sección y las fracturas. El inventario visual no se ha usado para adjudicar un subtipo formal cuando la ficha individual no permite hacerlo.",
 "parallels": "La comparación definitiva debe hacerse con el repertorio del grupo D de Rodríguez-Martín y con stylos de contextos arqueológicos comparables. En ausencia de un paralelo individual citado de forma verificable para el código, este estudio evita asignar equivalencias específicas.",
 "chron": "No se propone una fecha individual a partir de la carpeta. La datación requiere asociación estratigráfica, contexto de hallazgo y contraste con los paralelos tipológicos."
},
"E": {
 "intro": "El grupo E contiene objetos relacionados con el juego: fichas, peones y dados. La presencia de estas clases de objetos permite hablar de prácticas lúdicas, aunque la forma aislada no siempre permite identificar un juego preciso ni distinguir con seguridad entre juego de mesa, pieza de cómputo y reutilización secundaria.",
 "models": "E.I incluye fichas circulares lisas o decoradas, en ocasiones con numerales. E.II comprende peones hemisféricos, lisos o decorados mediante pequeños anillos y círculos incisos. E.IV agrupa dados, principalmente cúbicos, además de variantes paralelepípedas. Para los dados cúbicos, la disposición de los valores y la suma habitual de las caras opuestas son criterios útiles de identificación, pero deben comprobarse pieza a pieza.",
 "parallels": "Las fichas y los peones se comparan con Augusta Emerita, Torre Águila, Emporiae, Saguntum, Tarraco, Petavonium, Valeria, Carmo, Ilici y Torre de Palma. Para dados se citan Segobriga, Emporiae, Mérida, Tarraco, Pollentia, Baelo Claudia, Córdoba, Itálica, Termes, Lyon, Avenches, Aquincum y Caesarea Maritima.",
 "chron": "Los objetos lúdicos tienen una cronología amplia. Fichas y peones se encuentran en contextos de los siglos I-IV d.C., con ejemplos tardíos, mientras que los dados presentan una perduración prolongada. No son fósiles directores precisos sin contexto."
},
"F": {
 "intro": "La familia F reúne elementos relacionados con muebles, cierres y ensamblajes. Se han reconocido precintos o sellos, bisagras y una serie de componentes de mobiliario. La interpretación funcional debe apoyarse en la forma de inserción, perforaciones, ranuras, huellas de desgaste y relación con posibles partes de madera o metal.",
 "models": "F.II incluye dispositivos de cierre con un elemento fijo y otro móvil, ranuras y orificios para cordeles o varillas. F.III recoge elementos que pueden funcionar como charnelas o bisagras, entre ellos anillas y cilindros huecos. F.IV es una categoría más amplia de elementos de mobiliario; la ficha fotográfica y la relación con otros componentes son imprescindibles para precisar su función.",
 "parallels": "Las fichas comparativas nombran Gades, Augusta Emerita, Itálica, Emporiae y Lisboa para precintos y cierres; para bisagras se mencionan Emporiae, Baetulo, Tarraco, Barcino, Carteia, Mérida, Andelo, Termes, Delos, Cnosos, Lyon, Avenches, Augst y Magdalensberg.",
 "chron": "Los paralelos de bisagras abarcan cronologías extensas, desde fases prerromanas en algunos yacimientos hasta los siglos IV-V d.C. Los precintos o sellos citados aparecen entre los siglos I y III d.C. La datación individual de los códigos de Cástulo requiere sus contextos propios."
},
"H": {
 "intro": "H agrupa herramientas y elementos varios, especialmente mangos y remates de mangos. En los conjuntos de hueso trabajado, el mango suele identificarse por su morfología, por el alojamiento del elemento útil y por las huellas de fijación o desgaste. Un remate puede confundirse con una pieza completa si el objeto se conserva fragmentado.",
 "models": "H.II recoge mangos y H.III remates de mangos. La asignación individual debe revisar si existe canal, perforación, cavidad, extremo de inserción o transición entre el cuerpo y el remate. El código 0014014012 aparece en ambas subcarpetas, lo que debe tratarse como adscripción cruzada hasta revisar las imágenes y las fichas.",
 "parallels": "Los paralelos exactos dependen del tipo de mango y de su sistema de fijación. No se proponen equivalencias específicas cuando la documentación del código no las confirma.",
 "chron": "No se asigna cronología individual a esta familia a partir de la carpeta. Deben combinarse contexto estratigráfico, comparación morfológica y relación con los instrumentos o muebles a los que pudo pertenecer."
},
"I": {
 "intro": "El grupo I recoge objetos cuya función no se reconoce con seguridad. La categoría no es un defecto del inventario, sino una forma de preservar la incertidumbre. Pueden incluirse aros, elementos seudoanulares y piezas cuya forma no conserva rasgos diagnósticos suficientes.",
 "models": "Los aros pueden confundirse con anillos, aunque su tamaño, grosor y forma de fabricación han llevado a proponer usos de ensamblaje o aplicaciones textiles. Los elementos seudoanulares pueden ser recortes de diáfisis o anillas trabajadas. Las piezas de función desconocida deben describirse morfológicamente antes de formular una hipótesis funcional.",
 "parallels": "Las fichas comparativas citan Emporiae, Tarraco y Stobi para aros, y Augusta Emerita y Barcelona para elementos seudoanulares. Para piezas sin función conocida aparecen Conimbriga, Mérida, Clunia, Palencia y Bilbilis.",
 "chron": "La cronología suele ser genérica o depende del contexto del paralelo. No se restringe la fecha cuando la forma es demasiado simple o incompleta para una comparación precisa."
},
"J": {
 "intro": "El grupo J documenta fases de manufactura, especialmente desbastes y esbozos. Su valor reside en la tecnología: permite estudiar selección de materia prima, corte, preparación de soportes, abandono de piezas y gestión de los desechos. Estos restos no deben contarse como objetos acabados cuando se analiza la incidencia funcional.",
 "models": "J.I incluye desbastes, epífisis cortadas y diáfisis preparadas o serradas. J.II recoge esbozos cuya forma está parcialmente definida y cuyo trabajo se interrumpió por rotura, defecto o abandono. El análisis requiere registrar superficies de corte, orientación de las marcas, fracturas, regularidad del soporte y posibles huellas de herramientas.",
 "parallels": "Los repertorios citan talleres y contextos de Baetulo, Augusta Emerita, Regina, Tarraco, Córdoba, Valentia, Herrera de Pisuerga y Colonia Ulpia Traiana (Xanten). El paralelo tecnológico debe basarse en restos de trabajo comparables y no únicamente en la coincidencia formal.",
 "chron": "Las fichas señalan que desbastes y esbozos deben fecharse por su unidad estratigráfica y materiales asociados. Su presencia aislada no demuestra por sí sola la existencia de un taller."
}
}

def font_setup():
    regular = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
    bold = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
    italic = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Oblique.ttf"
    if Path(regular).exists():
        pdfmetrics.registerFont(TTFont("DV", regular))
        pdfmetrics.registerFont(TTFont("DV-Bold", bold))
        pdfmetrics.registerFont(TTFont("DV-Italic", italic))
        return "DV", "DV-Bold", "DV-Italic"
    return "Helvetica", "Helvetica-Bold", "Helvetica-Oblique"

FONT, FONT_B, FONT_I = font_setup()
styles = getSampleStyleSheet()
styles.add(ParagraphStyle(name="CoverTitle2", fontName=FONT_B, fontSize=25, leading=31,
    alignment=TA_CENTER, textColor=colors.HexColor("#24384A"), spaceAfter=12))
styles.add(ParagraphStyle(name="CoverSub2", fontName=FONT, fontSize=11, leading=16,
    alignment=TA_CENTER, textColor=colors.HexColor("#546777"), spaceAfter=6))
styles.add(ParagraphStyle(name="H1x", fontName=FONT_B, fontSize=17, leading=22,
    textColor=colors.HexColor("#24384A"), spaceBefore=9, spaceAfter=8, keepWithNext=True))
styles.add(ParagraphStyle(name="H2x", fontName=FONT_B, fontSize=12, leading=16,
    textColor=colors.HexColor("#31566D"), spaceBefore=9, spaceAfter=5, keepWithNext=True))
styles.add(ParagraphStyle(name="H3x", fontName=FONT_B, fontSize=9.4, leading=12,
    textColor=colors.HexColor("#3B5367"), spaceBefore=6, spaceAfter=3, keepWithNext=True))
styles.add(ParagraphStyle(name="Bodyx", fontName=FONT, fontSize=8.7, leading=12.5,
    alignment=TA_JUSTIFY, spaceAfter=5))
styles.add(ParagraphStyle(name="Smallx", fontName=FONT, fontSize=7.3, leading=9.5,
    alignment=TA_LEFT, spaceAfter=2))
styles.add(ParagraphStyle(name="Tinyx", fontName=FONT, fontSize=6.2, leading=7.7,
    alignment=TA_LEFT, spaceAfter=0))
styles.add(ParagraphStyle(name="Captionx", fontName=FONT, fontSize=6.6, leading=8,
    alignment=TA_CENTER, textColor=colors.HexColor("#394957"), spaceAfter=2))
styles.add(ParagraphStyle(name="Calloutx", fontName=FONT, fontSize=8.4, leading=12,
    textColor=colors.HexColor("#263746"), backColor=colors.HexColor("#F0F4F6"),
    borderColor=colors.HexColor("#D5DEE5"), borderWidth=0.45, borderPadding=7,
    spaceBefore=3, spaceAfter=7))
styles.add(ParagraphStyle(name="TableHeadx", fontName=FONT_B, fontSize=7, leading=8.5,
    textColor=colors.white, alignment=TA_LEFT))
styles.add(ParagraphStyle(name="TableCellx", fontName=FONT, fontSize=6.8, leading=8.5,
    alignment=TA_LEFT))
styles.add(ParagraphStyle(name="TableCellSmallx", fontName=FONT, fontSize=5.9, leading=7.4,
    alignment=TA_LEFT))
def esc(x):
    return str(x).replace("&","&amp;").replace("<","&lt;").replace(">","&gt;")
def P(txt, style="Bodyx"):
    return Paragraph(txt, styles[style])

def family_for_path(path):
    try:
        rel = path.relative_to(ROOT)
        top = rel.parts[0]
    except Exception:
        top = path.parts[0] if path.parts else ""
    m = re.match(r"^([A-J])\s*[-.]", top)
    if not m:
        return "?"
    return m.group(1)

def code_for_image(p):
    stem = p.stem
    # Primary inventory code before .1, .2, -vista, etc.
    m = re.match(r"^(\d{4,12})(?:[._ -].*)?$", stem)
    if m:
        return m.group(1)
    # Special IDs are retained separately, never coerced to numeric inventory.
    m = re.match(r"^((?:DJ|ALJ)\d{4,12})(?:[._ -].*)?$", stem, re.I)
    if m:
        return m.group(1).upper()
    if stem.lower().startswith("harp"):
        return "SUP-HARPOCRATES"
    if "whatsapp" in stem.lower():
        return "SUP-WHATSAPP"
    if "avestruz" in stem.lower() or "huevo" in stem.lower():
        return "SUP-AVESTRUZ"
    return "SUP-" + re.sub(r"[^A-Za-z0-9_-]+","_",stem)[:48]

def human_path(p):
    try:
        return str(p.relative_to(ROOT))
    except Exception:
        return str(p)

def collect_images():
    if not ROOT.exists():
        raise FileNotFoundError(f"No existe la carpeta esperada: {ROOT}")
    files = [p for p in ROOT.rglob("*") if p.is_file() and p.suffix.lower() in (".jpg",".jpeg",".png",".tif",".tiff",".bmp",".webp")]
    grouped = defaultdict(list)
    for p in files:
        grouped[code_for_image(p)].append(p)
    return files, grouped

def prepare_image(path):
    # Resize and recompress while preserving aspect ratio; no destructive changes to repository.
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    out = WORK / (digest + ".jpg")
    if out.exists():
        return out, digest
    with PILImage.open(path) as im:
        im = ImageOps.exif_transpose(im)
        if im.mode not in ("RGB","L"):
            bg = PILImage.new("RGB", im.size, "white")
            if "A" in im.getbands():
                bg.paste(im, mask=im.getchannel("A"))
                im = bg
            else:
                im = im.convert("RGB")
        if im.mode == "L":
            im = im.convert("RGB")
        im.thumbnail((1100, 850), PILImage.Resampling.LANCZOS)
        im.save(out, "JPEG", quality=73, optimize=True, progressive=True)
    return out, digest

def extract_doc_notes():
    notes = defaultdict(list)
    docs = [p for p in ROOT.rglob("*.docx") if not p.name.startswith("~$") and p.stat().st_size < 8_000_000]
    for path in docs:
        try:
            d = Document(str(path))
            lines = []
            for para in d.paragraphs:
                s = re.sub(r"\s+"," ",para.text or "").strip()
                if s:
                    lines.append(s)
            for table in d.tables:
                for row in table.rows:
                    s = " | ".join(re.sub(r"\s+"," ",c.text or "").strip() for c in row.cells)
                    if s.strip(" |"):
                        lines.append(s)
            for i, line in enumerate(lines):
                ids = re.findall(r"(?<!\d)(\d{8,12})(?!\d)", line)
                for code in set(ids):
                    context = " ".join(lines[max(0,i-1):min(len(lines),i+3)])
                    context = re.sub(r"\s+"," ",context).strip()
                    if context and context not in notes[code]:
                        notes[code].append(f"{path.name}: {context[:1100]}")
        except Exception as e:
            print(f"Warning: no se pudo leer {path.name}: {e}", file=sys.stderr)
    return notes, len(docs)

class FamilyBars(Flowable):
    def __init__(self, rows, width=485, height=190):
        super().__init__()
        self.rows=rows; self.width=width; self.height=height
    def draw(self):
        c=self.canv
        maxv=max([x[1] for x in self.rows] + [1])
        left=155; barw=self.width-left-48; rowh=17
        y=self.height-17
        for label,val in self.rows:
            c.setFont(FONT,6.7); c.setFillColor(colors.HexColor("#283D4E"))
            c.drawRightString(left-7,y+2,label[:32])
            c.setFillColor(colors.HexColor("#E4EAF0")); c.rect(left,y,barw,8,stroke=0,fill=1)
            c.setFillColor(colors.HexColor("#426985")); c.rect(left,y,barw*val/maxv,8,stroke=0,fill=1)
            c.setFont(FONT_B,6.7); c.setFillColor(colors.HexColor("#24384A"))
            c.drawString(left+barw+7,y+1,str(val))
            y-=rowh

def page_footer(canvas, doc):
    canvas.saveState()
    w,h=A4
    if doc.page>1:
        canvas.setFillColor(colors.HexColor("#60717F"))
        canvas.setFont(FONT_B,7)
        canvas.drawString(17*mm,h-11*mm,"CÁSTULO | INDUSTRIA ÓSEA")
        canvas.setStrokeColor(colors.HexColor("#D4DCE2")); canvas.setLineWidth(.45)
        canvas.line(17*mm,h-13*mm,w-17*mm,h-13*mm)
    canvas.setStrokeColor(colors.HexColor("#D4DCE2")); canvas.line(17*mm,13*mm,w-17*mm,13*mm)
    canvas.setFont(FONT,6.5); canvas.setFillColor(colors.HexColor("#6B7785"))
    canvas.drawString(17*mm,8*mm,"Estudio tipológico e inventario fotográfico | versión de trabajo, 29-10-2026")
    canvas.drawRightString(w-17*mm,8*mm,str(doc.page))
    canvas.restoreState()

def main():
    paths, grouped = collect_images()
    notes, n_docs = extract_doc_notes()
    # Gather IDs by primary family; cross-classification remains visible.
    family_codes = defaultdict(set)
    image_info = {}
    figure_number = 0
    for code, originals in grouped.items():
        # Exact duplicate photo files are merged for display, while provenance paths remain listed.
        by_hash = {}
        for p in sorted(originals, key=lambda z: str(z).lower()):
            fam = family_for_path(p)
            if fam != "?":
                family_codes[fam].add(code)
            try:
                jpg, digest = prepare_image(p)
            except Exception as e:
                print(f"Warning: imagen omitida {p}: {e}", file=sys.stderr)
                continue
            if digest not in by_hash:
                by_hash[digest] = {"path": p, "jpg": jpg, "sources": [], "families": set()}
            by_hash[digest]["sources"].append(human_path(p))
            if fam != "?":
                by_hash[digest]["families"].add(fam)
        ims = list(by_hash.values())
        ims.sort(key=lambda z: str(z["path"]).lower())
        image_info[code] = ims

    numeric_codes = sorted([c for c in image_info if c.isdigit()], key=lambda z:(int(z),len(z),z))
    supplementary_codes = sorted([c for c in image_info if not c.isdigit()])
    total_numeric = len(numeric_codes)
    cross = {c: sorted({family_for_path(x) for x in grouped[c] if family_for_path(x)!="?"}) for c in numeric_codes}
    cross_codes = [c for c,v in cross.items() if len(v)>1]
    # Keep output details in log for final verification.
    family_rows = [(k, len(family_codes[k])) for k in "ABCDEFGHIJ" if family_codes[k]]
    print("PDF inventory summary:", {"image_files":len(paths), "numeric_codes":total_numeric,
          "supplementary_codes":len(supplementary_codes), "distinct_photo_digests":sum(len(x) for x in image_info.values()),
          "cross_family_codes":len(cross_codes), "families":family_rows, "docx_texts_read":n_docs})

    doc=SimpleDocTemplate(str(OUT), pagesize=A4, rightMargin=17*mm, leftMargin=17*mm,
        topMargin=20*mm, bottomMargin=18*mm, title="La industria ósea de Cástulo: estudio extenso e ilustrado",
        author="Estudio arqueológico de materiales", subject="Clasificación tipológica, catálogo fotográfico y valoración cuantitativa")
    story=[]
    story += [Spacer(1,28*mm), P("LA INDUSTRIA ÓSEA<br/>DE CÁSTULO","CoverTitle2"),
        P("Estudio extenso de clasificación tipológica, tecnología, cronología, paralelos e incidencia cuantitativa","CoverSub2"),
        Spacer(1,8*mm), HRFlowable(width="60%",thickness=1,color=colors.HexColor("#9CAAB7"),hAlign="CENTER"),
        Spacer(1,9*mm), P("Catálogo ilustrado elaborado a partir de la carpeta <i>Clasificación-Huesos-de-Castulo-(7-10-26)</i> del repositorio <b>germanroma-alt/Castulo</b>.","CoverSub2"),
        Spacer(1,13*mm), P("Marco tipológico principal: F. Germán Rodríguez-Martín, <i>La industria ósea en la Hispania romana</i>, Archaeopress, 2024.","CoverSub2"),
        Spacer(1,18*mm), P(f"{total_numeric} códigos numéricos reconocidos en las imágenes de inventario<br/>{len(paths)} archivos de imagen revisados; fotografías repetidas agrupadas por código y duplicados exactos detectados por huella digital<br/>Fecha de elaboración: 29 de octubre de 2026","CoverSub2"),
        Spacer(1,23*mm), P("Documento de investigación y trabajo. Las atribuciones tipológicas y cronológicas se expresan según el grado de certeza permitido por las fichas conservadas.","CoverSub2"), PageBreak()]
    # Executive overview
    story += [P("Resumen", "H1x"),
        P(f"El directorio fotográfico permite reconocer <b>{total_numeric} códigos numéricos</b> de pieza, además de {len(supplementary_codes)} identificadores complementarios no normalizados. Se han revisado {len(paths)} archivos de imagen; estos no equivalen al número de objetos, puesto que un mismo código puede contar con distintas vistas y aparecer en más de una subcarpeta. Para preservar la integridad de la colección se ha organizado un único registro por código, agrupando bajo él las fotografías asociadas y conservando su procedencia de carpeta."),
        P("El conjunto muestra variedad funcional: instrumentos de trabajo textil, elementos de adorno y toilette, accesorios de vestimenta, escritura, objetos lúdicos, elementos de mobiliario, mangos y piezas de uso incierto, así como restos de fabricación. El predominio estadístico de una familia debe entenderse como una propiedad del inventario digitalizado, no como un indicador directo de la frecuencia de uso en toda la ciudad."),
        P("La relación de cada figura con su objeto se establece a partir del código que figura en el nombre de archivo. Cuando un mismo código aparece bajo más de una familia, la adscripción doble queda señalada y no se crea una segunda pieza ficticia. Las fotografías sin código numérico se apartan al final como documentación complementaria.", "Calloutx"),
        P("Guía de lectura", "H2x"),
        P("La primera parte desarrolla contexto y metodología; la segunda presenta la distribución cuantitativa y la interpretación por familias; la tercera organiza un catálogo fotográfico en el que cada código tiene una entrada, fotografías numeradas, procedencia del archivo y una ficha textual recuperada de los documentos disponibles cuando existe coincidencia explícita. El anexo final recoge identificadores especiales y advertencias sobre las adscripciones cruzadas."),
        P("Contenido", "H2x")]
    toc = [
      "1. Objeto y alcance del estudio", "2. Contexto arqueológico e historia de la ciudad",
      "3. Fuentes documentales y metodología", "4. Criterios de clasificación y grados de certeza",
      "5. Recuento y distribución de las familias", "6. Estudio tipológico por familias A-J",
      "7. Materia prima, técnicas y cadena operativa", "8. Cronología y uso de los paralelos",
      "9. Incidencia del conjunto e interpretación", "10. Conclusiones y líneas de trabajo",
      "Catálogo fotográfico principal por código", "Anexos: adscripciones cruzadas e imágenes suplementarias",
    ]
    story.append(Table([[P("• "+esc(s),"Smallx")] for s in toc], colWidths=[168*mm],
       style=TableStyle([("LEFTPADDING",(0,0),(-1,-1),5),("RIGHTPADDING",(0,0),(-1,-1),5),
                         ("TOPPADDING",(0,0),(-1,-1),2),("BOTTOMPADDING",(0,0),(-1,-1),2)])))
    # Main study
    story += [P("1. Objeto y alcance del estudio","H1x"),
      P("El objetivo es transformar el archivo de clasificación de Cástulo en un instrumento de análisis arqueológico: ordenar los objetos por familias funcionales; hacer explícita la correspondencia entre cada código y las imágenes conservadas; sintetizar la información tipológica de las fichas; confrontar modelos y paralelos; y valorar la incidencia cuantitativa del conjunto. El inventario incluye objetos acabados, elementos fragmentarios, posibles componentes de otros objetos y desechos de manufactura, clases que no deben mezclarse sin advertencia."),
      P("El código de inventario se utiliza como unidad de recuento básica. Las vistas numeradas con sufijos como .1, .2 o .3 se consideran fotografías de la misma pieza mientras el prefijo del identificador sea idéntico. Si el código aparece en dos familias, se conserva la adscripción cruzada y se evita duplicar el objeto en el total global. La coincidencia de código se trata como indicio administrativo; los casos con dudas funcionales se mantienen abiertos a revisión."),
      P("El estudio no sustituye la ficha arqueológica original ni una revisión directa de las piezas. Las fotografías permiten estudiar formas y acabados, pero la determinación anatómica, taxonómica, huellas microscópicas, medidas exactas, estado de conservación y contexto de excavación requieren los registros específicos o el examen del original."),
      P("2. Contexto arqueológico e historia de la ciudad","H1x"),
      P("Cástulo, en el alto Guadalquivir, fue uno de los principales centros urbanos de la región durante la Antigüedad. Su larga trayectoria ibérica y romana, el papel de la ciudad en las redes de circulación y su transformación urbana explican que el conjunto de cultura material deba interpretarse en una perspectiva diacrónica. Sin embargo, el carácter urbano del yacimiento no convierte a todos los materiales del inventario en productos locales: los objetos óseos pueden haber sido fabricados en talleres de la ciudad, adquiridos mediante intercambio, reutilizados o desplazados en depósitos secundarios."),
      P("Para el estudio de industria ósea, el contexto local es relevante en tres niveles: la función del objeto, el momento de utilización y deposición, y la posible producción. La función se propone por comparación morfológica; la fecha depende de la unidad estratigráfica y sus asociaciones; y la fabricación local necesita indicios tecnológicos contextualizados, tales como desechos, soportes, herramientas y piezas en distintas fases de manufactura. Ninguno de estos niveles puede sustituir al otro."),
      P("La carpeta revisada no contiene un subdirectorio G de objetos relacionados con el mundo militar. Esta ausencia debe expresarse como un resultado documental, no como prueba de inexistencia arqueológica. La ausencia de imágenes o una representación muy baja de una categoría puede deberse a la selección de la colección, a las áreas excavadas, a la conservación diferencial o a decisiones de organización del archivo."),
      P("3. Fuentes documentales y metodología","H1x"),
      P("La fuente primaria es el árbol de archivos de la carpeta denominada Clasificación-Huesos-de-Castulo-(7-10-26), con subcarpetas funcionales A, B, C, D, E, F, H, I y J y documentos de clasificación. Se han revisado los nombres de las fotografías, la estructura de las subcarpetas y los documentos Word cuyo tamaño permite una extracción textual automatizada. La extracción de texto se utiliza como localizador de descripciones, no como sustituto de la lectura de cada ficha completa."),
      P("Los archivos fotográficos se agrupan por código. Las copias digitales idénticas se identifican mediante huella SHA-256 y se muestran una sola vez en la entrada principal; sus rutas de procedencia quedan consignadas para que sea posible reconstruir la organización original. Las imágenes distintas con el mismo código se conservan como vistas separadas. La numeración de figuras es correlativa y se genera a partir de los archivos realmente insertados."),
      P("Para evitar extrapolaciones, se distinguen cuatro clases de dato: (1) código y ruta comprobados; (2) clasificación de carpeta; (3) descripción textual recuperada de una ficha; y (4) interpretación propuesta por comparación. Si no aparece una descripción literal para un código, la ficha del catálogo lo declara y ofrece únicamente la clasificación documental general. La tipología y cronología de un modelo comparativo no se presentan como datación individual de la pieza."),
      P("4. Criterios de clasificación y grados de certeza","H1x"),
      P("El marco tipológico sigue la estructura A-J empleada en la obra de Rodríguez-Martín. A agrupa objetos textiles; B, objetos de vida cotidiana, adorno y toilette; C, vestido; D, escritura; E, juego y divertimento; F, mobiliario; G, mundo militar; H, herramientas y elementos varios; I, objetos sin función conocida; y J, fases de elaboración. En el inventario digital consultado no hay fotografías en una carpeta G separada."),
      P("Dentro de cada familia, la asignación al tipo y subtipo se realiza solo cuando el documento o la imagen ofrece rasgos diagnósticos suficientes. En alfileres, por ejemplo, el remate y la sección del vástago pueden ser determinantes; en agujas, la forma del ojal; en dados, la geometría del cuerpo y la distribución de puntos; en bisagras, la relación entre anilla, cilindro y eje; en restos de manufactura, la orientación y naturaleza de los cortes. Un nombre de carpeta es una primera adscripción, no siempre un subtipo definitivo."),
      P("Las adscripciones cruzadas no se resuelven mediante una elección arbitraria. Si un código aparece en dos familias, se mantiene una sola entrada y se enumeran sus ubicaciones. Esto preserva el recuento de piezas y, a la vez, permite auditar el sistema de clasificación. En casos como un código localizado en B.IX/Punzones, se deja constancia de la anomalía del árbol y se remite a revisión de imágenes y ficha."),
      P("5. Recuento y distribución de las familias","H1x"),
      P(f"El recuento automático identifica {total_numeric} códigos numéricos únicos en el árbol fotográfico. Hay {len(cross_codes)} códigos que aparecen en más de una familia de primer nivel. La cifra exacta de archivos de imagen ({len(paths)}) incluye vistas diferentes y fotografías con nombres especiales, por lo que no debe confundirse con el número de objetos. La tabla siguiente expresa códigos distintos presentes en cada familia; por la existencia de adscripciones cruzadas, los porcentajes por familia no son excluyentes."),
    ]
    family_table = [[P("Grupo","TableHeadx"),P("Familia","TableHeadx"),P("Códigos distintos","TableHeadx"),P("% sobre códigos numéricos","TableHeadx")]]
    for code,name in FAMILY_NAMES.items():
        n=len(family_codes[code])
        if n:
            family_table.append([P(code,"TableCellx"),P(esc(name),"TableCellx"),P(str(n),"TableCellx"),P(f"{100*n/max(total_numeric,1):.1f}%","TableCellx")])
    ft=Table(family_table,colWidths=[16*mm,73*mm,32*mm,47*mm],repeatRows=1)
    ft.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#34566E")),
        ("GRID",(0,0),(-1,-1),.35,colors.HexColor("#D1DAE1")),("VALIGN",(0,0),(-1,-1),"TOP"),
        ("LEFTPADDING",(0,0),(-1,-1),5),("RIGHTPADDING",(0,0),(-1,-1),5),
        ("TOPPADDING",(0,0),(-1,-1),4),("BOTTOMPADDING",(0,0),(-1,-1),4),
        ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#F5F7F9")])]))
    story += [ft,Spacer(1,5),FamilyBars([(f"{k} - {FAMILY_NAMES[k]}",len(family_codes[k])) for k in "ABCDEFGHIJ" if family_codes[k]],width=485,height=190),
        P("La familia B reúne el repertorio de uso personal y toilette y suele ser la más representada; su amplitud taxonómica explica parte de esa incidencia. E agrupa prácticas lúdicas, A evidencia actividades textiles y F refleja componentes de muebles o cierres. J es especialmente significativo para la tecnología, aunque los restos de trabajo deben separarse de los objetos funcionales cuando se comparan categorías de uso."),
        P("Los números son una fotografía documental del catálogo digital en la fecha de revisión. No constituyen un censo de todos los materiales óseos excavados en Cástulo ni permiten inferir directamente el peso económico de cada actividad. Antes de utilizar porcentajes como indicador de comportamiento social es necesario conocer los criterios de selección, procedencia, cronología, conservación y publicación de las piezas.")]
    # Family studies
    story += [PageBreak(),P("6. Estudio tipológico por familias","H1x")]
    for code in "ABCDEFGHIJ":
        if code not in FAMILY_TEXT:
            continue
        if not family_codes[code]:
            continue
        f=FAMILY_TEXT[code]
        n=len(family_codes[code])
        story += [P(f"{code}. {FAMILY_NAMES[code]}","H2x"),
          P(f"<b>Incidencia documental.</b> La familia contiene {n} códigos numéricos distintos en las imágenes revisadas ({100*n/max(total_numeric,1):.1f}% del total numérico). La tasa es de presencia y puede solaparse con otras familias cuando el mismo código aparece en más de una carpeta."),
          P("<b>Caracterización funcional y morfológica.</b> "+esc(f["intro"])),
          P("<b>Modelos y rasgos diagnósticos.</b> "+esc(f["models"])),
          P("<b>Paralelos.</b> "+esc(f["parallels"])),
          P("<b>Cronología.</b> "+esc(f["chron"])),
          P("<b>Valor interpretativo.</b> El grupo debe ponerse en relación con el resto del repertorio, el estado de conservación y la información contextual. La función propuesta no equivale a la demostración de un uso exclusivo; la cronología comparativa no sustituye a la datación arqueológica del ejemplar.")]
    story += [P("6.1. El mundo militar (grupo G)","H2x"),
      P("La estructura general de la clasificación contempla un grupo G para objetos del mundo militar, pero no se ha identificado una carpeta G con fotografías dentro del directorio consultado. Este apartado se conserva para hacer explícito el vacío documental. No se incorporan piezas hipotéticas ni se concluye que no existieran objetos de esa clase en la ciudad.")]
    story += [P("7. Materia prima, técnicas y cadena operativa","H1x"),
      P("La categoría industria ósea se usa como etiqueta de conjunto, pero no permite asignar sin análisis la especie animal ni distinguir siempre hueso, asta y marfil. En este estudio, la materia prima se considera determinada solo cuando lo señala la ficha o cuando el registro de colección ofrece una atribución comprobable. No se presupone que todos los objetos de una misma forma estén hechos del mismo material."),
      P("Las operaciones tecnológicas que deben analizarse en las fotografías y fichas comprenden selección de soporte, corte, serrado, desbaste, raspado, regularización de superficies, perforación, torneado, incisión y pulido. La forma de los cortes, la regularidad de los vástagos, la alineación de perforaciones y el tratamiento de los extremos pueden indicar el grado de estandarización y la secuencia de trabajo. No todas estas operaciones son visibles en fotografía general y las huellas finas requieren observación directa o macrofotografía."),
      P("La familia J ofrece una clave para reconstruir la cadena operativa. Los desbastes pueden reflejar la preparación inicial de una materia prima; los esbozos muestran formas parcialmente definidas que no llegaron a completarse; los objetos acabados documentan el resultado. La presencia de estos tres estados en un mismo contexto reforzaría una hipótesis de producción, pero la co-ocurrencia de materiales en una colección sin proveniencia estratigráfica no basta para demostrar un taller."),
      P("Una futura revisión debería registrar por pieza: elemento anatómico o soporte, especie cuando pueda determinarse, dimensiones máximas, sección, huellas de herramienta, presencia de perforaciones y su secuencia, acabado, fracturas antiguas o recientes, uso/desgaste, restauración, contexto arqueológico y paralelos. Las medidas se deben cotejar con el código exacto, no trasladar desde una tabla general de colección sin conciliación.")]
    story += [P("8. Cronología y uso de los paralelos","H1x"),
      P("Las fichas tipológicas proporcionan horquillas de cronología de los modelos, pero los paralelos no son todos equivalentes: algunos proceden de conjuntos bien fechados, otros de contextos amplios y otros se citan por semejanza formal. La fuerza cronológica de una comparación depende de su contexto, de la continuidad del tipo y de la semejanza real de los rasgos diagnósticos."),
      P("Para cada objeto se deben distinguir tres fechas posibles: la del modelo en el conjunto comparativo; la del contexto arqueológico en que apareció el ejemplar de Cástulo; y la del momento de fabricación o uso, que puede preceder al depósito. En caso de ausencia de estratigrafía, solo puede formularse una propuesta tipológica y debe expresarse como tal."),
      P("Los grupos A y B incluyen formas de larga duración, particularmente agujas, alfileres y objetos de adorno. E también reúne tipos que circulan durante siglos. Esos objetos suelen tener valor funcional pero escaso poder para precisar una fecha estrecha. Las piezas lúdicas cúbicas y los elementos simples de mobiliario son ejemplos de formas relativamente perdurables. Los modelos raros, por su parte, pueden ser útiles para comparar redes de circulación, pero su rareza no garantiza una datación más precisa."),
      P("Por ello, el catálogo individual incorpora una nota temporal de prudencia. Cuando la ficha documental incluye una cronología contextual o un paralelo explícito, se cita; cuando solo hay una adscripción de carpeta, no se inventa una fecha concreta. La ausencia de fecha por pieza es un resultado metodológicamente correcto, no una omisión que deba rellenarse con una horquilla genérica.")]
    story += [P("9. Incidencia del conjunto e interpretación","H1x"),
      P("El repertorio combina actividades de cuidado personal, adorno, prácticas lúdicas, producción textil, escritura, componentes de muebles y operaciones de fabricación. Si se confirma el predominio de la familia B en un inventario más completo y controlado, la explicación inicial debe considerar tanto el interés arqueológico de esos objetos como la amplitud de categorías que reúne. No conviene reducir el patrón a una lectura directa de las costumbres de la población."),
      P("El conjunto textil tiene interés porque incluye instrumentos de hilado y costura y placas de telar, y puede ayudar a estudiar tanto el trabajo doméstico como la fabricación de productos para intercambio. Sin embargo, el número de instrumentos por sí solo no cuantifica la producción, ni separa actividad profesional de actividad doméstica. Para ello son necesarios contexto, asociaciones, huellas de uso y comparaciones con otros materiales."),
      P("La presencia de fichas, peones y dados prueba la disponibilidad de objetos que pudieron intervenir en juegos; no identifica en exclusiva reglas, edad o género de sus usuarios. Los objetos de toilette y adorno remiten a prácticas de cuidado personal e imagen, pero no deben leerse como marcadores sociales automáticos sin atender a materia prima, calidad, decoración y contexto."),
      P("Los elementos de mobiliario y mangos muestran que las materias duras animales podían integrarse en objetos compuestos junto a madera, metal o materiales orgánicos perecederos. En estos casos, la conservación diferencial puede dejar visible el componente óseo pero hacer desaparecer el objeto al que pertenecía. La interpretación debe reconstruir esa relación con prudencia."),
      P("Los restos de desbaste y esbozos son importantes para discutir manufactura, pero la prueba de un taller exige una concentración contextual de residuos, herramientas y productos en diferentes fases. Una colección acumulada procedente de múltiples intervenciones puede reunir piezas funcionalmente relacionadas sin que hayan formado parte de un mismo proceso productivo.")]
    story += [P("10. Conclusiones y líneas de trabajo","H1x"),
      P(f"1. La carpeta fotográfica permite reconocer {total_numeric} códigos numéricos de pieza y {len(supplementary_codes)} identificadores suplementarios no normalizados. El total de archivos de imagen es mayor porque el mismo objeto tiene varias vistas y algunos códigos figuran en más de una familia."),
      P("2. La numeración de figuras se ha generado a partir de fotografías efectivamente asociadas al código; las vistas de un objeto se agrupan bajo una sola entrada, y las copias idénticas se reducen sin borrar las rutas de procedencia."),
      P("3. Las familias tipológicas permiten comparar el repertorio de Cástulo con la obra de Rodríguez-Martín, pero solo se adjudican subtipos concretos cuando existen rasgos diagnósticos o una ficha explícita."),
      P("4. Los paralelos peninsulares y extrapeninsulares son útiles para situar los modelos; la cronología de cada pieza debe contrastarse con el contexto arqueológico y no puede deducirse automáticamente del modelo."),
      P("5. La incidencia cuantitativa describe el inventario digitalizado, no toda la industria ósea de Cástulo ni, por sí sola, la intensidad social o económica de una actividad."),
      P("6. El grupo J ofrece la vía más directa para desarrollar una investigación tecnológica, siempre que se relacione con procedencia, huellas de trabajo, soportes y herramientas."),
      P("7. La siguiente fase editorial debería completar para cada objeto la materia prima, medidas, conservación, unidad estratigráfica y referencia bibliográfica individual, a partir del inventario arqueológico completo.")]
    # Catalogue
    story += [PageBreak(),P("11. Catálogo fotográfico principal por código","H1x"),
      P("Las entradas se ordenan por código de inventario. Cada objeto tiene un único registro aunque aparezca en varias familias. Debajo de cada código se muestran las fotografías localizadas, cada una con número de figura correlativo, nombre del archivo de origen y carpeta de procedencia. Las fotografías idénticas entre rutas distintas se muestran una vez y las rutas quedan indicadas en el pie."),
      P("La descripción textual que se incluye por pieza procede de las fichas Word cuando se ha podido localizar el código. Si no existe un fragmento documental enlazable de manera segura, el catálogo distingue ese vacío y mantiene la descripción al nivel de la clasificación de la carpeta. Las fechas son tipológicas y contextuales solo cuando la documentación las aporta.")]
    fig_no=0
    code_to_primary={}
    for code in numeric_codes:
        ims=image_info.get(code,[])
        fams=cross.get(code,[])
        famlabels=[f"{f} - {FAMILY_NAMES.get(f,'familia no identificada')}" for f in fams]
        folder_paths=sorted({human_path(p).split(os.sep)[0] for p in grouped[code]})
        subpaths=sorted({human_path(p).rsplit(os.sep,1)[0] for p in grouped[code]})
        code_notes=notes.get(code,[])
        # One object entry starts with metadata.
        title=f"Código {esc(code)}"
        if fams:
            title += " | " + esc("; ".join(famlabels))
        block=[P(title,"H3x"),
            P("<b>Adscripción documental.</b> "+esc("; ".join(subpaths[:8]))+("…" if len(subpaths)>8 else ""),"Smallx")]
        if len(fams)>1:
            block.append(P("<b>Nota de clasificación cruzada.</b> Este código aparece en más de una familia de primer nivel. Se cuenta como un objeto y se conservan las adscripciones para su revisión morfofuncional.","Smallx"))
        if code_notes:
            # Keep only up to 2 distinct short notes per object to maintain readability.
            used=[]
            for n in code_notes:
                if n not in used:
                    used.append(n)
                if len(used)>=2: break
            block.append(P("<b>Información textual recuperada.</b> "+esc(" || ".join(used)),"Smallx"))
        else:
            block.append(P("<b>Descripción individual.</b> No se ha localizado en las fichas Word extraídas un pasaje que cite expresamente este código. La identificación se mantiene en el nivel de la carpeta y la imagen original; no se atribuyen medidas, materia prima taxonómica o subtipo que no estén documentados.","Smallx"))
        family_data = [FAMILY_TEXT[f] for f in fams if f in FAMILY_TEXT]
        if family_data:
            block.append(P("<b>Marco interpretativo.</b> "+esc(family_data[0]["chron"]),"Smallx"))
        if ims:
            # images arranged two across
            cells=[]
            for im in ims:
                fig_no += 1
                code_to_primary.setdefault(code, fig_no)
                im["fig"] = fig_no
                im_obj=Image(str(im["jpg"]), width=77*mm, height=51*mm, kind="proportional")
                source_paths = im["sources"]
                cap=f"Figura {fig_no}. Código {code}. Vista {esc(im['path'].stem)}."
                if len(source_paths)>1:
                    cap += f" Archivo duplicado en {len(source_paths)} rutas; se muestra una sola copia digital."
                cap += f"<br/>Origen: {esc(' | '.join(source_paths[:2]))}"
                if len(source_paths)>2: cap += " (y otras rutas)"
                cells.append([im_obj,P(cap,"Captionx")])
            # table two columns; pad last row
            rows=[]
            for i in range(0,len(cells),2):
                left=cells[i]; right=cells[i+1] if i+1<len(cells) else [P("","Smallx"),P("","Captionx")]
                rows.append([left,right])
            img_table=Table(rows,colWidths=[81*mm,81*mm],hAlign="LEFT")
            img_table.setStyle(TableStyle([
                ("VALIGN",(0,0),(-1,-1),"TOP"),
                ("LEFTPADDING",(0,0),(-1,-1),2),("RIGHTPADDING",(0,0),(-1,-1),2),
                ("TOPPADDING",(0,0),(-1,-1),3),("BOTTOMPADDING",(0,0),(-1,-1),4),
                ("BOX",(0,0),(-1,-1),0.3,colors.HexColor("#E0E5EA")),
                ("INNERGRID",(0,0),(-1,-1),0.25,colors.HexColor("#E0E5EA")),
            ]))
            block.append(img_table)
        else:
            block.append(P("No se ha podido preparar una imagen legible para este código; revisar el archivo original.","Smallx"))
        story.append(KeepTogether(block))
        story.append(Spacer(1,5))
    # Supplemental gallery
    if supplementary_codes:
        story += [PageBreak(),P("12. Identificadores suplementarios e imágenes sin código numérico","H1x"),
            P(f"Se han reconocido {len(supplementary_codes)} identificadores o nombres de imagen que no forman parte del recuento numérico principal. Pueden incluir códigos alfanuméricos, nombres descriptivos, vistas sin una clave normalizada o documentación relacionada con materiales no incluidos en la industria ósea estricta. Se muestran como documentación auxiliar y no se suman a los {total_numeric} códigos numéricos.")]
        for code in supplementary_codes:
            ims=image_info.get(code,[])
            block=[P(f"Identificador suplementario: {esc(code)}","H3x")]
            srcs=sorted({human_path(p) for p in grouped[code]})
            block.append(P("<b>Procedencia.</b> "+esc("; ".join(srcs[:8]))+("…" if len(srcs)>8 else ""),"Smallx"))
            if code=="SUP-AVESTRUZ":
                block.append(P("Este archivo se considera material complementario: una cáscara o huevo de avestruz no se incorpora al recuento de industria ósea en sentido estricto.","Smallx"))
            cells=[]
            for im in ims:
                fig_no+=1
                im["fig"]=fig_no
                img=Image(str(im["jpg"]),width=77*mm,height=51*mm,kind="proportional")
                cap=f"Figura {fig_no}. Identificador {esc(code)}. Archivo {esc(im['path'].name)}.<br/>Origen: {esc(' | '.join(im['sources'][:2]))}"
                cells.append([img,P(cap,"Captionx")])
            rows=[]
            for i in range(0,len(cells),2):
                rows.append([cells[i],cells[i+1] if i+1<len(cells) else [P("","Smallx"),P("","Captionx")]])
            if rows:
                tab=Table(rows,colWidths=[81*mm,81*mm])
                tab.setStyle(TableStyle([("VALIGN",(0,0),(-1,-1),"TOP"),("GRID",(0,0),(-1,-1),.25,colors.HexColor("#E0E5EA")),("LEFTPADDING",(0,0),(-1,-1),2),("RIGHTPADDING",(0,0),(-1,-1),2)]))
                block.append(tab)
            story.append(KeepTogether(block)); story.append(Spacer(1,5))
    # Crossover schedule
    story += [PageBreak(),P("Anexo I. Códigos con adscripción a varias familias","H1x"),
      P("Los códigos siguientes figuran en más de una familia de primer nivel. Se presentan como una relación de control; la interpretación debe apoyarse en las fotografías y la ficha. La doble clasificación puede deberse a una adscripción funcional diferente, a una reorganización de carpetas o a una pieza que comparte rasgos con más de un tipo.")]
    cd=[[P("Código","TableHeadx"),P("Familias detectadas","TableHeadx"),P("N.º de fotos distintas","TableHeadx")]]
    for code in cross_codes:
        nf=len(image_info.get(code,[]))
        cd.append([P(esc(code),"TableCellx"),P(esc(", ".join(f"{x}: {FAMILY_NAMES[x]}" for x in cross[code])),"TableCellx"),P(str(nf),"TableCellx")])
    ctab=LongTable(cd,colWidths=[30*mm,110*mm,27*mm],repeatRows=1)
    ctab.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#34566E")),("GRID",(0,0),(-1,-1),.3,colors.HexColor("#D1DAE1")),("VALIGN",(0,0),(-1,-1),"TOP"),("LEFTPADDING",(0,0),(-1,-1),4),("RIGHTPADDING",(0,0),(-1,-1),4),("TOPPADDING",(0,0),(-1,-1),3),("BOTTOMPADDING",(0,0),(-1,-1),3),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#F5F7F9")])]))
    story.append(ctab)
    story += [P("Anexo II. Limitaciones y tareas de verificación","H1x"),
      P("1. Conciliar los códigos del registro fotográfico con el inventario arqueológico y las unidades estratigráficas. 2. Revisar las fotografías originales para asociar cada código con medidas y descripción morfológica. 3. Completar la determinación de materia prima cuando exista base anatómica o analítica. 4. Asegurar que la cronología tipológica no se confunda con la cronología contextual. 5. Normalizar subcarpetas con nomenclaturas históricas o inconsistentes. 6. Resolver las adscripciones cruzadas y mantener una tabla de cambios. 7. Verificar la bibliografía completa de cada paralelo en la obra tipológica de referencia."),
      P("Anexo III. Bibliografía y fuentes","H1x"),
      P("RODRÍGUEZ-MARTÍN, F. Germán (2024): <i>La industria ósea en la Hispania romana</i>. Archaeopress Roman Archaeology 110. Oxford. DOI: 10.32028/9781803276656."),
      P("Fuente del inventario: repositorio GitHub <i>germanroma-alt/Castulo</i>, carpeta <i>Clasificación-Huesos-de-Castulo-(7-10-26)</i>. Documentación de referencia consultada según los ficheros disponibles: fichas de acus, telarcillos, acus crinalis, espátulas de toilette, cucharillas de perfume, strigiles, pulseras, colgantes, broches, stylos, fichas de juego, peones, dados, precintos/sellos, bisagras, mobiliario, mangos, aros, piezas sin función conocida, desbastes y esbozos."),
      P("Nota final: este volumen es una versión ampliada e ilustrada del archivo digital. No sustituye un catálogo de excavación con contexto, dimensiones, conservación y bibliografía individual plenamente cotejados. Cuando falta ese dato, la carencia se señala en lugar de completarse con una inferencia no verificable.")]
    doc.build(story,onFirstPage=page_footer,onLaterPages=page_footer)
    print("Generated",OUT,OUT.stat().st_size,"bytes, figures",fig_no)
    print("Family counts",[(k,len(family_codes[k])) for k in "ABCDEFGHIJ"])
    print("Catalog entries",len(numeric_codes),"Supplementary entries",len(supplementary_codes))

if __name__=="__main__":
    main()
